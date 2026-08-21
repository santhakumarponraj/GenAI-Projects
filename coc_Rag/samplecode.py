import numpy as np
import os
import pdfplumber
from sentence_transformers import SentenceTransformer
import docx
from ollama import chat



cocFile = fr"C:\Users\SanthakumarPonraj\Downloads\All complaint\All complaint\CoC_JST Sales America Inc_ASSHSSH28K51_canadapfas.docx"


def read_docx(file_path):
    doc = docx.Document(file_path)

    pages = []
    current_page = []

    for paragraph in doc.paragraphs:
        # Check for explicit page break
        xml = paragraph._element.xml

        if 'w:type="page"' in xml:
            pages.append("\n".join(current_page))
            current_page = []

        current_page.append(paragraph.text)

    # Add remaining text
    if current_page:
        pages.append("\n".join(current_page))

    # Add page numbers
    text = ""

    for page_number, page_text in enumerate(pages, start=1):
        text += f"\n[Page Number {page_number}]\n"
        text += page_text + "\n"

    return text

    
def read_pdf(file_path):
    text = ""

    with pdfplumber.open(file_path) as pdf:
        for pagenumber,page in enumerate(pdf.pages):
            text += f"\n [Page Number {pagenumber+1}] \n"

            text += page.extract_text()

    return text


def chunk_text(pdf_text, chunk_size=500, overlap=100):
    chunks = []

    start = 0
    step = chunk_size - overlap

    while start < len(pdf_text):
        end = start + chunk_size

        chunk = pdf_text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        start += step

    return chunks


def embedding_process(embedding_model, chunks):
    embeddings = []

    for chunk in chunks:
        chunk_embedding = embedding_model.encode(
            chunk,
            convert_to_numpy=True,
            normalize_embeddings=True
        )

        embeddings.append(chunk_embedding)

    return np.array(embeddings)

def vector_Memory_Store(chunks,chunk_embedding):
    vector_store = []

    for chunk,embedding in zip(chunks,chunk_embedding):

        vector_store.append({
            "chunk" : chunk,
            'embedding' : embedding})

    return vector_store

def Search_Document(question,vector_store,embedding_model, top_k = 5):

    query_embedding= embedding_model.encode(
        question,
        convert_to_numpy= True,
        normalize_embeddings= True
    )

    embeddings = np.array([item["embedding"] for item in vector_store])

    scores = np.dot(embeddings,query_embedding)

    top_indices = np.argsort(scores)[::-1][:top_k]

    result = []

    for index in top_indices:

        result.append({
            "chunk" : vector_store[index]["chunk"],
            "score" : float(scores[index])
        })

    return result

def build_context(result):

    context = ""

    for i,result in enumerate(result):
        context += f"""--- Document Chunk { i+1 } ---
        {result["chunk"]}"""

    return context


def generate_answer(question, context):

    prompt = f"""
    You are an expert Compliance Document Analysis Assistant.

    Your task is to analyze supplier compliance documents such as:
    - Certificate of Compliance (CoC)
    - Declaration of Conformity (DoC)
    - Material Declaration
    - Environmental Compliance Reports
    - Regulatory Compliance Certificates

    Use ONLY the information available in the provided context.

    RULES:

    1. Identify the regulation(s) covered in the document.
    Examples:
    - California Proposition 65 (CA Prop 65)
    - RoHS
    - REACH
    - TSCA
    - PFAS
    - Conflict Minerals
    - POPs
    - ELV
    - Halogen Free

    2. Determine the compliance status.

    Compliance Status Values:
    - Compliant
    - Non-Compliant
    - Exempted
    - Not Determined
    - Unknown

    3. Extract supporting evidence directly from the document.

    4. If multiple regulations are mentioned, provide status for each regulation separately.

    5. Never infer or assume information not explicitly stated in the document.

    6. If the document does not contain enough information, answer:
    'I could not find the answer from the given document.'

    7. Return the result in JSON format.

    CONTEXT:
    {context}

    QUESTION:
    {question}

    RESPONSE FORMAT:

    {{
    "document_type": "",
    "regulations": [
        {{
        "regulation_name": "",
        "compliance_status": "",
        "evidence": ""
        }}
    ],
    "part_number": "",
    "certificate_number": "",
    "issue_date": "",
    "summary": ""
    }}
    """

    response = chat(
        model="qwen2.5vl:3b",
        messages=[{
            "role": "user",
            "content": prompt}])

    return response.message.content




if __name__ == "__main__":

    text = read_docx(cocFile)

    chunks = chunk_text(text)

    embedding_model = SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")

    chunk_embedding = embedding_process(embedding_model,chunks)

    vector_store = vector_Memory_Store(chunks,chunk_embedding)

    question = "Is this product compliant?"

    result = Search_Document(question,vector_store,embedding_model)


    context = build_context(result)

    answer = generate_answer(question,context)

    print(answer)
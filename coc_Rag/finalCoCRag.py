import numpy as np
import pdfplumber
from sentence_transformers import SentenceTransformer
import docx
from ollama import chat
import torch
import os


def read_docx(file_path):
    doc = docx.Document(file_path)

    pages = []
    current_page = []

    for paragraph in doc.paragraphs:
        xml = paragraph._element.xml
        if 'w:type="page"' in xml:
            pages.append("\n".join(current_page))
            current_page = []
        current_page.append(paragraph.text)

    if current_page:
        pages.append("\n".join(current_page))

    # Use list + join instead of += in a loop (avoids O(n^2) string copies)
    parts = []
    for page_number, page_text in enumerate(pages, start=1):
        parts.append(f"\n[Page Number {page_number}]\n{page_text}\n")

    return "".join(parts)


def read_pdf(file_path):
    parts = []
    with pdfplumber.open(file_path) as pdf:
        for pagenumber, page in enumerate(pdf.pages):
            parts.append(f"\n [Page Number {pagenumber+1}] \n")
            parts.append(page.extract_text() or "")
    return "".join(parts)


def chunk_text(pdf_text, chunk_size=500, overlap=100):
    chunks = []
    start = 0
    step = chunk_size - overlap
    n = len(pdf_text)

    while start < n:
        chunk = pdf_text[start:start + chunk_size].strip()
        if chunk:
            chunks.append(chunk)
        start += step

    return chunks


def embedding_process(embedding_model, chunks, batch_size=32):
    # Batch-encode everything in one call instead of looping chunk by chunk.
    # This lets the model process multiple chunks per forward pass.
    embeddings = embedding_model.encode(
        chunks,
        batch_size=batch_size,
        convert_to_numpy=True,
        normalize_embeddings=True,
        show_progress_bar=False,
    )
    return embeddings


def vector_Memory_Store(chunks, chunk_embeddings):
    # Keep chunks and the embedding matrix separate so we don't rebuild
    # the matrix on every search call.
    return {
        "chunks": chunks,
        "embeddings": chunk_embeddings,  # already a single (N, D) array
    }


def Search_Document(question, vector_store, embedding_model, top_k=5):
    query_embedding = embedding_model.encode(
        question,
        convert_to_numpy=True,
        normalize_embeddings=True,
    )

    embeddings = vector_store["embeddings"]
    scores = embeddings @ query_embedding

    # Clamp top_k so it never exceeds the number of chunks available
    k = min(top_k, len(scores))

    top_indices = np.argpartition(scores, -k)[-k:]
    top_indices = top_indices[np.argsort(scores[top_indices])[::-1]]

    return [
        {"chunk": vector_store["chunks"][i], "score": float(scores[i])}
        for i in top_indices
    ]


def build_context(results):
    parts = []
    for i, r in enumerate(results):
        parts.append(f"--- Document Chunk {i+1} ---\n{r['chunk']}\n")
    return "".join(parts)


def generate_answer(question, context):
    prompt = f"""
    You are an expert Compliance Document Analysis Assistant.

    Your task is to analyze supplier compliance documents such as:
    - Certificate of Compliance (CoC)
    - Declaration of Conformity (DoC)
    - Material Declaration
    - Environmental Compliance Reports
    - Regulatory Compliance Certificates

    Use ONLY the information available in the provided context.

    RULES:

    1. Identify the regulation(s) covered in the document.
    Examples:
    - California Proposition 65 (CA Prop 65)
    - RoHS
    - REACH
    - TSCA
    - PFAS
    - Conflict Minerals
    - POPs
    - ELV
    - Halogen Free

    2. Determine the compliance status.

    Compliance Status Values:
    - Compliant
    - Non-Compliant
    - Exempted
    - Not Determined
    - Unknown

    3. Extract supporting evidence directly from the document.

    4. If multiple regulations are mentioned, provide status for each regulation separately.

    5. Never infer or assume information not explicitly stated in the document.

    6. If the document does not contain enough information, answer:
    'I could not find the answer from the given document.'

    7. Return the result in JSON format.

    CONTEXT:
    {context}

    QUESTION:
    {question}

    RESPONSE FORMAT:

    {{
    "document_type": "",
    "regulations": [
        {{
        "regulation_name": "",
        "compliance_status": "",
        "evidence": ""
        }}
    ],
    "part_number": "",
    "certificate_number": "",
    "issue_date": "",
    "summary": ""
    }}
    """

    response = chat(
        model="qwen2.5vl:3b",
        messages=[{"role": "user", "content": prompt}],
    )

    return response.message.content


import os

if __name__ == "__main__":

    while True:

        cocFile = input("Enter the FilePath : ==> 'exit','cancel','end' <== : ").strip()

        if cocFile.lower() in ['exit', 'cancel', 'end']:
            break

        cocFile = os.path.abspath(cocFile)

        print("Path received:", cocFile)
        print("File exists:", os.path.exists(cocFile))

        if cocFile.lower().endswith(".docx"):
            text = read_docx(cocFile)

        elif cocFile.lower().endswith(".pdf"):
            text = read_pdf(cocFile)

        else:
            print("==== Format should be .pdf or .docx ====")
            continue

        chunks = chunk_text(text)

        device = "cuda" if torch.cuda.is_available() else "cpu"
        embedding_model = SentenceTransformer(
            "sentence-transformers/all-MiniLM-L6-v2",
            device=device,
        )

        chunk_embeddings = embedding_process(embedding_model, chunks)
        vector_store = vector_Memory_Store(chunks, chunk_embeddings)

        # question = "Is this product compliant?"

        while True:

            question = input("Enter the User Question : ==> 'exit','cancel','end' <== : ")

            if question.lower() in ['exit','cancel','end']:
                break

            result = Search_Document(question, vector_store, embedding_model, top_k=3)

            context = build_context(result)
            answer = generate_answer(question, context)

            print(answer)